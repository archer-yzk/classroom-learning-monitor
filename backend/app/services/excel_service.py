"""使用 openpyxl 生成历史学情 Excel 文件。"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.app.core.time import utc_naive_to_china
from backend.app.services.history_service import HistoryDatabaseRow


HISTORY_HEADERS = [
    "序号",
    "课程名称",
    "班级",
    "课堂主题",
    "上课时间",
    "统计开始时间",
    "统计结束时间",
    "统计间隔（分钟）",
    "提交学生人数",
    "1分人数",
    "2分人数",
    "3分人数",
    "4分人数",
    "5分人数",
    "6分人数",
    "平均分",
    "方差",
]


def describe_query_conditions(
    course_name: str | None,
    start_time: datetime | None,
    end_time: datetime | None,
) -> str:
    """生成人能读懂的导出条件说明。"""

    course_text = course_name.strip() if course_name and course_name.strip() else "全部课程"
    local_start_time = utc_naive_to_china(start_time)
    local_end_time = utc_naive_to_china(end_time)
    start_text = (
        local_start_time.strftime("%Y-%m-%d %H:%M")
        if local_start_time
        else "不限"
    )
    end_text = (
        local_end_time.strftime("%Y-%m-%d %H:%M")
        if local_end_time
        else "不限"
    )
    return f"查询条件：课程={course_text}；上课时间={start_text} 至 {end_text}"


def build_history_excel(
    rows: list[HistoryDatabaseRow],
    teacher_name: str,
    course_name: str | None,
    start_time: datetime | None,
    end_time: datetime | None,
    exported_at: datetime,
) -> BytesIO:
    """把历史快照写入内存中的 Excel 工作簿并返回文件流。"""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "学情历史"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A5"

    # 标题和导出说明保持简洁，表头从第 4 行开始。
    worksheet["A1"] = "课堂学情历史统计"
    worksheet["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="1F2937")
    worksheet["A2"] = (
        f"授课教师：{teacher_name}    "
        f"导出时间：{exported_at.strftime('%Y-%m-%d %H:%M:%S')}"
    )
    worksheet["A2"].font = Font(name="Microsoft YaHei", size=10, italic=True, color="667085")
    worksheet["A3"] = describe_query_conditions(course_name, start_time, end_time)
    worksheet["A3"].font = Font(name="Microsoft YaHei", size=10, italic=True, color="667085")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    header_border = Border(right=Side(style="thin", color="FFFFFF"))

    for column_index, header in enumerate(HISTORY_HEADERS, start=1):
        cell = worksheet.cell(row=4, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    date_columns = {5, 6, 7}
    text_columns = {2, 3, 4}
    light_fill = PatternFill("solid", fgColor="F4F7FB")
    bottom_border = Border(bottom=Side(style="thin", color="D9E2F0"))

    for row_index, (snapshot, classroom, course) in enumerate(rows, start=5):
        values = [
            row_index - 4,
            course.name,
            classroom.class_name,
            classroom.topic,
            utc_naive_to_china(classroom.start_time),
            utc_naive_to_china(snapshot.period_start),
            utc_naive_to_china(snapshot.period_end),
            snapshot.interval_minutes,
            snapshot.total_students,
            snapshot.score_1_count,
            snapshot.score_2_count,
            snapshot.score_3_count,
            snapshot.score_4_count,
            snapshot.score_5_count,
            snapshot.score_6_count,
            snapshot.average_score,
            snapshot.variance,
        ]
        worksheet.append(values)

        for column_index in range(1, len(HISTORY_HEADERS) + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            cell.font = Font(name="Microsoft YaHei", size=10, color="1F2937")
            cell.border = bottom_border
            cell.alignment = Alignment(
                horizontal="left" if column_index in text_columns else "right",
                vertical="center",
                wrap_text=column_index in text_columns,
            )
            if column_index in date_columns:
                cell.number_format = "yyyy-mm-dd hh:mm"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif column_index in {16, 17}:
                cell.number_format = "0.00"

            if row_index % 2 == 0:
                cell.fill = light_fill

        worksheet.row_dimensions[row_index].height = 30

    last_row = 4 + len(rows)
    worksheet.row_dimensions[1].height = 26
    worksheet.row_dimensions[4].height = 34

    column_widths = [8, 20, 18, 26, 19, 19, 19, 17, 15, 11, 11, 11, 11, 11, 11, 12, 12]
    for column_index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    worksheet.print_title_rows = "1:4"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.auto_filter.ref = f"A4:Q{last_row}"

    workbook.properties.title = "课堂学情历史统计"
    workbook.properties.subject = "教师历史学情查询导出"
    workbook.properties.creator = "课堂学情监控管理系统"

    file_stream = BytesIO()
    workbook.save(file_stream)
    file_stream.seek(0)
    return file_stream
